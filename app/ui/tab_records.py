from __future__ import annotations
import datetime
import re

import openpyxl
from PyQt5.QtCore import Qt, pyqtSignal
from PyQt5.QtWidgets import (
    QFileDialog, QHBoxLayout, QHeaderView, QLabel, QMessageBox,
    QPushButton, QSplitter, QTableWidget, QTableWidgetItem,
    QVBoxLayout, QWidget,
)

from app.data.models import ClassData
from app.data.repository import Repository
from app.logic.statistics import calc_class_stats


class RecordsTab(QWidget):
    records_changed = pyqtSignal()

    def __init__(self, class_data: ClassData, repo: Repository, parent=None):
        super().__init__(parent)
        self._class_data = class_data
        self._repo = repo
        self._init_ui()
        self._refresh()

    def set_class_data(self, class_data: ClassData) -> None:
        self._class_data = class_data
        self._refresh()

    def _init_ui(self):
        layout = QVBoxLayout(self)

        splitter = QSplitter(Qt.Vertical)

        stats_widget = QWidget()
        stats_layout = QVBoxLayout(stats_widget)
        stats_layout.setContentsMargins(0, 4, 0, 0)
        stats_layout.addWidget(QLabel("学生统计"))
        self._stats_table = QTableWidget(0, 4)
        self._stats_table.setHorizontalHeaderLabels(["学号", "姓名", "评分次数", "平均分"])
        self._stats_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self._stats_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self._stats_table.setSelectionBehavior(QTableWidget.SelectRows)
        stats_layout.addWidget(self._stats_table)
        splitter.addWidget(stats_widget)

        history_widget = QWidget()
        history_layout = QVBoxLayout(history_widget)
        history_layout.setContentsMargins(0, 4, 0, 0)
        history_layout.addWidget(QLabel("点名历史"))
        self._history_table = QTableWidget(0, 4)
        self._history_table.setHorizontalHeaderLabels(["学号", "姓名", "时间", "分数"])
        self._history_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self._history_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self._history_table.setSelectionBehavior(QTableWidget.SelectRows)
        history_layout.addWidget(self._history_table)
        splitter.addWidget(history_widget)

        layout.addWidget(splitter)

        btn_row = QHBoxLayout()
        del_btn = QPushButton("删除选中历史记录")
        del_btn.clicked.connect(self._delete_record)
        export_btn = QPushButton("导出本班成绩单")
        export_btn.clicked.connect(self._export_single)
        export_all_btn = QPushButton("导出全部班级")
        export_all_btn.clicked.connect(self._export_all)
        for btn in (del_btn, export_btn, export_all_btn):
            btn_row.addWidget(btn)
        btn_row.addStretch()
        layout.addLayout(btn_row)

    def _refresh(self):
        stats_by_id = {s["student_id"]: s for s in calc_class_stats(self._class_data.records)}

        self._stats_table.setRowCount(0)
        for student in self._class_data.students:
            row = self._stats_table.rowCount()
            self._stats_table.insertRow(row)
            self._stats_table.setItem(row, 0, QTableWidgetItem(student.id))
            self._stats_table.setItem(row, 1, QTableWidgetItem(student.name))
            s = stats_by_id.get(student.id)
            if s:
                self._stats_table.setItem(row, 2, QTableWidgetItem(str(s["count"])))
                self._stats_table.setItem(row, 3, QTableWidgetItem(str(s["avg_score"])))
            else:
                self._stats_table.setItem(row, 2, QTableWidgetItem("0"))
                self._stats_table.setItem(row, 3, QTableWidgetItem("-"))

        self._history_table.setRowCount(0)
        for r in reversed(self._class_data.records):
            row = self._history_table.rowCount()
            self._history_table.insertRow(row)
            self._history_table.setItem(row, 0, QTableWidgetItem(r.student_id))
            self._history_table.setItem(row, 1, QTableWidgetItem(r.name))
            self._history_table.setItem(row, 2, QTableWidgetItem(r.time))
            self._history_table.setItem(row, 3, QTableWidgetItem(str(r.score)))

    def _delete_record(self):
        row = self._history_table.currentRow()
        if row < 0:
            QMessageBox.warning(self, "提示", "请先选中一条历史记录。")
            return
        real_idx = len(self._class_data.records) - 1 - row
        try:
            with self._repo.edit_class(self._class_data):
                del self._class_data.records[real_idx]
        except Exception as e:
            QMessageBox.critical(self, "删除失败", str(e))
            return
        self._refresh()
        self.records_changed.emit()

    def _export_single(self):
        default_name = f"{self._class_data.name}成绩单_{datetime.date.today()}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "导出成绩单", default_name, "Excel Files (*.xlsx)")
        if not path:
            return
        try:
            wb = openpyxl.Workbook()
            self._write_class_sheet(wb.active, self._class_data.name, self._class_data.students, self._class_data.records)
            wb.save(path)
        except Exception as e:
            QMessageBox.critical(self, "导出失败", str(e))
            return
        QMessageBox.information(self, "成功", "成绩单已导出。")

    def _export_all(self):
        default_name = f"全部班级成绩单_{datetime.date.today()}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "导出全部班级", default_name, "Excel Files (*.xlsx)")
        if not path:
            return
        try:
            wb = openpyxl.Workbook()
            all_classes = self._repo.data.classes if self._repo.data else []
            for index, cls in enumerate(all_classes):
                ws = wb.active if index == 0 else wb.create_sheet()
                self._write_class_sheet(ws, cls.name, cls.students, cls.records)
            wb.save(path)
        except Exception as e:
            QMessageBox.critical(self, "导出失败", str(e))
            return
        QMessageBox.information(self, "成功", "全部班级成绩单已导出。")

    @staticmethod
    def _write_class_sheet(ws, class_name: str, students, records) -> None:
        title = re.sub(r'[\\/*?:\[\]]', '_', class_name).strip("'")[:31] or "班级"
        existing = {s.title.lower() for s in ws.parent.worksheets if s is not ws}
        base = title
        suffix = 1
        while title.lower() in existing:
            tail = f"_{suffix}"
            title = base[:31 - len(tail)] + tail
            suffix += 1
        ws.title = title
        ws.append(["学号", "姓名", "评分次数", "平均分"])
        stats_by_id = {s["student_id"]: s for s in calc_class_stats(records)}
        for student in students:
            s = stats_by_id.get(student.id)
            if s:
                ws.append([student.id, student.name, s["count"], s["avg_score"]])
            else:
                ws.append([student.id, student.name, 0, "-"])
