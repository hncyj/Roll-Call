import json

import openpyxl
import pytest
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import QFileDialog, QMessageBox, QInputDialog

from app.data.models import ClassData, Student
from app.data.repository import Repository
from app.ui.main_window import MainWindow
from app.ui.dialogs.class_dialog import DeleteClassDialog
from app.ui.dialogs.settings_dialog import SettingsDialog
from app.ui.tab_records import RecordsTab


@pytest.fixture
def window(qapp, messages, tmp_path):
    repo = Repository(str(tmp_path / "data.json"))
    data = repo.load()
    data.classes = [
        ClassData("a", "A", [Student("001", "Alice"), Student("002", "Bob"), Student("003", "Carol")]),
        ClassData("b", "B", [Student("001", "Other Alice")]),
    ]
    repo.save()
    widget = MainWindow(repo)
    yield widget
    widget._roll_call_tab.cancel_selection()
    widget.close()
    widget.deleteLater()
    qapp.processEvents()


def finish_roll(tab):
    for _ in tab._animation._INTERVALS:
        tab._animation._tick()


def test_random_round_persists_without_scores_and_never_repeats(window):
    tab = window._roll_call_tab
    ids = []
    for _ in range(3):
        tab._random_roll()
        assert tab._animation.is_animating
        finish_roll(tab)
        ids.append(tab._selected_student.id)
        persisted = Repository(window._repo.data_path).load().classes[0]
        assert persisted.called_ids == ids
    assert len(set(ids)) == 3
    assert tab._class_data.records == []
    tab._random_roll()
    assert not tab._animation.is_animating
    assert tab._engine.pick_random() is None


def test_restart_after_custom_path_keeps_round_without_score(window, tmp_path):
    window._repo.set_data_path(str(tmp_path / "external" / "roster.json"))
    tab = window._roll_call_tab
    tab._random_roll()
    finish_roll(tab)
    sid = tab._selected_student.id
    restarted = MainWindow()
    try:
        assert restarted._repo.data_path == window._repo.data_path
        assert sid not in [s.id for s in restarted._roll_call_tab._engine.get_uncalled()]
        assert restarted._data.classes[0].records == []
    finally:
        restarted.close()


def test_switch_class_cancels_animation_and_rejects_old_result(window):
    tab = window._roll_call_tab
    tab._random_roll()
    old_target = tab._animation._target
    window._class_combo.setCurrentIndex(1)
    assert not tab._animation.is_animating
    finish_roll(tab)
    tab._on_animation_done(old_target)
    tab._save_record()
    assert tab._selected_student is None
    assert window._data.classes[0].called_ids == []
    assert window._data.classes[1].records == []
    assert window._data.classes[1].called_ids == []
    assert tab._random_btn.isEnabled()


def test_same_result_can_only_be_scored_once(window):
    tab = window._roll_call_tab
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    tab._save_record()
    tab._save_record()
    assert len(tab._class_data.records) == 1
    assert not tab._save_btn.isEnabled()


def test_cannot_save_previous_student_during_new_animation(window):
    tab = window._roll_call_tab
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    tab._random_roll()
    tab._save_record()
    assert not tab._save_btn.isEnabled()
    assert tab._class_data.records == []


def test_manual_call_also_excludes_already_called_students(window, messages):
    tab = window._roll_call_tab
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    tab._save_record()
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    assert tab._selected_student is None
    assert tab._class_data.called_ids == ["001"]
    assert messages[-1][1] == "本轮已点名"


def test_reset_requires_confirmation_and_persists(window, monkeypatch):
    tab = window._roll_call_tab
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    tab._reset_called()
    assert tab._class_data.called_ids == ["001"]
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    tab._reset_called()
    assert tab._class_data.called_ids == []
    assert tab._selected_student is None
    assert Repository(window._repo.data_path).load().classes[0].called_ids == []


def test_failed_call_does_not_mark_success(window, monkeypatch, messages):
    tab = window._roll_call_tab
    tab._random_roll()
    def fail():
        raise OSError("disk full")
    monkeypatch.setattr(window._repo, "save", fail)
    finish_roll(tab)
    assert tab._selected_student is None
    assert tab._class_data.called_ids == []
    assert not tab._save_btn.isEnabled()
    assert messages[-1][0] == "critical"


def test_failed_score_can_be_retried_without_duplicate(window, monkeypatch):
    tab = window._roll_call_tab
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    original_save = window._repo.save
    def fail():
        raise OSError("disk full")
    monkeypatch.setattr(window._repo, "save", fail)
    tab._save_record()
    assert tab._class_data.records == []
    assert tab._class_data.called_ids == ["001"]
    assert tab._selected_student is not None
    monkeypatch.setattr(window._repo, "save", original_save)
    tab._save_record()
    assert len(tab._class_data.records) == 1


def test_delete_last_class_rebinds_every_tab_to_saved_default(window, monkeypatch):
    monkeypatch.setattr(DeleteClassDialog, "exec_", lambda self: DeleteClassDialog.Accepted)
    window._delete_class()
    window._roll_call_tab._random_roll()
    window._delete_class()
    cls = window._data.classes[0]
    assert len(window._data.classes) == 1
    assert cls.name == "默认班级"
    assert cls.students == []
    assert not window._roll_call_tab._animation.is_animating
    for tab in (window._students_tab, window._records_tab, window._roll_call_tab):
        assert tab._class_data is cls
    assert Repository(window._repo.data_path).load().classes[0].id == cls.id


def test_deleting_student_refreshes_manual_selection_by_id(window, monkeypatch):
    tab = window._roll_call_tab
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    window._students_tab._table.setCurrentCell(0, 0)
    window._students_tab._delete_student()
    assert tab._student_list.count() == 2
    assert tab._student_list.item(0).data(Qt.UserRole) == "002"
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    assert tab._selected_student.name == "Bob"


def test_edit_id_updates_score_and_round_in_ui(window, monkeypatch):
    tab = window._roll_call_tab
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    tab._save_record()
    answers = iter([("101", True), ("Alice renamed", True)])
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: next(answers))
    window._students_tab._table.setCurrentCell(0, 0)
    window._students_tab._edit_student()
    assert tab._class_data.called_ids == ["101"]
    assert tab._class_data.records[0].student_id == "101"
    assert window._records_tab._stats_table.item(0, 2).text() == "1"
    assert tab._student_list.item(0).data(Qt.UserRole) == "101"


@pytest.mark.parametrize("duplicate", [False, True])
def test_import_preserves_round_and_rejects_duplicate_ids(window, tmp_path, monkeypatch, duplicate):
    tab = window._roll_call_tab
    tab._student_list.setCurrentRow(0)
    tab._manual_roll()
    book = openpyxl.Workbook()
    book.active.append(["学号", "姓名"])
    book.active.append(["001", "Alice"])
    book.active.append(["001" if duplicate else "004", "New student"])
    path = tmp_path / "roster.xlsx"
    book.save(path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(path), ""))
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    window._students_tab._import_excel()
    assert tab._class_data.called_ids == ["001"]
    assert len(tab._class_data.students) == (3 if duplicate else 2)
    assert "001" not in [s.id for s in tab._engine.get_uncalled()]
    assert Repository(window._repo.data_path).load().classes[0].called_ids == ["001"]


def test_open_existing_file_rebinds_pages_without_overwriting(window, tmp_path, monkeypatch):
    path = tmp_path / "saved.json"
    path.write_text(json.dumps({"version": 2, "classes": [ClassData("saved", "Saved").to_dict()]}))
    before = path.read_bytes()
    monkeypatch.setattr(SettingsDialog, "exec_", lambda self: SettingsDialog.Accepted)
    monkeypatch.setattr(SettingsDialog, "selected_path", lambda self: str(path))
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    window._open_settings()
    assert window._data.classes[0].id == "saved"
    assert window._roll_call_tab._class_data is window._data.classes[0]
    assert path.read_bytes() == before


def test_export_handles_invalid_and_duplicate_long_class_names(tmp_path):
    wb = openpyxl.Workbook()
    name = "Class:/?[]" * 6
    RecordsTab._write_class_sheet(wb.active, name, [Student("001", "Alice")], [])
    RecordsTab._write_class_sheet(wb.create_sheet(), name, [], [])
    assert len(set(wb.sheetnames)) == 2
    assert all(len(title) <= 31 for title in wb.sheetnames)
    path = tmp_path / "export.xlsx"
    wb.save(path)
    restored = openpyxl.load_workbook(path)
    assert restored.worksheets[0]["A2"].value == "001"
    restored.close()
