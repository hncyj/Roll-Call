import openpyxl
import pytest
from PyQt5.QtWidgets import QFileDialog

from app.data.models import ClassData, Record, Student
from app.data.repository import Repository
from app.ui.tab_records import RecordsTab


@pytest.fixture
def grade_class():
    return ClassData("a", "一班", [
        Student("001", "张三"), Student("002", "张三"), Student("003", "王五"),
    ], [
        Record("001", "张三", "2026-09-03 09:00:00", 100),
        Record("002", "张三", "2026-09-01 09:00:00", 90),
        Record("001", "张三", "2026-09-01 09:00:00", 0),
        Record("001", "张三", "2026-09-02 09:00:00", 85),
        Record("002", "张三", "2026-09-02 09:00:00", 80),
    ])


def test_export_each_score_in_time_order_with_count_and_average(tmp_path, grade_class):
    original = list(grade_class.records)
    wb = openpyxl.Workbook()
    RecordsTab._write_class_sheet(wb.active, grade_class.name, grade_class.students, grade_class.records)
    path = tmp_path / "grades.xlsx"
    wb.save(path)
    saved = openpyxl.load_workbook(path, data_only=True)
    assert list(saved.active.values) == [
        ("学号", "姓名", "评分次数", "第 1 次评分", "第 2 次评分", "第 3 次评分", "平均分"),
        ("001", "张三", 3, 0, 85, 100, 61.7),
        ("002", "张三", 2, 90, 80, None, 85),
        ("003", "王五", 0, None, None, None, "-"),
    ]
    saved.close()
    assert grade_class.records == original


def test_same_timestamp_keeps_original_record_order(grade_class):
    grade_class.records = [Record("001", "张三", "2026-09-01 09:00:00", s) for s in (82, 91, 75)]
    wb = openpyxl.Workbook()
    RecordsTab._write_class_sheet(wb.active, grade_class.name, grade_class.students, grade_class.records)
    assert list(wb.active.values)[1][3:6] == (82, 91, 75)


@pytest.mark.parametrize("students", [[], [Student("001", "张三")]])
def test_no_scores_does_not_create_phantom_score_columns(students):
    wb = openpyxl.Workbook()
    RecordsTab._write_class_sheet(wb.active, "一班", students, [])
    assert list(wb.active.values)[0] == ("学号", "姓名", "评分次数", "平均分")
    if students:
        assert list(wb.active.values)[1] == ("001", "张三", 0, "-")


def test_deleted_students_do_not_add_unused_score_columns(grade_class):
    grade_class.records += [Record("deleted", "已删除", "2026-09-01", 70)] * 10
    wb = openpyxl.Workbook()
    RecordsTab._write_class_sheet(wb.active, grade_class.name, grade_class.students, grade_class.records)
    assert wb.active.max_column == 7
    assert wb.active.max_row == 4


@pytest.mark.parametrize("export_all", [False, True])
def test_export_buttons_write_score_columns(qapp, messages, tmp_path, monkeypatch, grade_class, export_all):
    repo = Repository(str(tmp_path / "data.json"))
    data = repo.load()
    data.classes = [grade_class, ClassData("b", "二班", [Student("001", "李四")], [
        Record("001", "李四", "2026-09-04 10:00:00", 76),
    ])]
    repo.save()
    tab = RecordsTab(grade_class, repo)
    path = tmp_path / "export.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(path), ""))
    if export_all:
        tab._export_all()
    else:
        tab._export_single()
    saved = openpyxl.load_workbook(path, data_only=True)
    try:
        assert list(saved["一班"].values)[1] == ("001", "张三", 3, 0, 85, 100, 61.7)
        assert saved.sheetnames == (["一班", "二班"] if export_all else ["一班"])
        if export_all:
            assert list(saved["二班"].values) == [
                ("学号", "姓名", "评分次数", "第 1 次评分", "平均分"),
                ("001", "李四", 1, 76, 76),
            ]
        assert messages[-1][0] == "information"
    finally:
        saved.close()
        tab.close()
